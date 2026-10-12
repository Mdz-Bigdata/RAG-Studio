import os
import uuid
import shutil
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, BackgroundTasks, status
from sqlalchemy.orm import Session
from qdrant_client.http import models as rest_models
from app.config import settings
from app.repository import get_db, SQLiteRepository, qdrant_repository
from app.model import DocumentModel, ChunkModel
from app.schema import DocumentResponse, DocumentEditRequest
from app.service import parser_service, embedding_service, unweaver_service

router = APIRouter()

# 存储上传文件的本地文件夹
UPLOAD_DIR = "./uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

def process_document_task(doc_id: str, kb_id: str, file_path: str):
    """
    异步后台任务：调用 MinerU/pdfplumber 解析 PDF，进行文本切片、向量化并入库，同时提取 UnWeaver 实体。
    """
    # 建立独立的数据库 Session
    from app.repository import SessionLocal
    db = SessionLocal()
    repo = SQLiteRepository(db)
    
    try:
        # 1. 更新状态为解析中
        repo.update_document_status(doc_id, "PROCESSING")
        
        # 2. 调用解析服务解析 PDF
        markdown_content = parser_service.parse_pdf(file_path)
        
        # 3. 保存解析得到的 Markdown 并做向量化与实体关系编织
        index_document_markdown(db, repo, doc_id, kb_id, markdown_content)
        
    except Exception as e:
        print(f"Error processing document {doc_id}: {e}")
        repo.update_document_status(doc_id, "FAILED", error_msg=str(e))
    finally:
        db.close()

def index_document_markdown(db: Session, repo: SQLiteRepository, doc_id: str, kb_id: str, markdown_content: str):
    """
    对 Markdown 进行切片、向量化、入库，并提取实体关系 (共享给上传与手工编辑更新流程)
    """
    # 1. 保存 Markdown 文本到 Document 表中
    repo.update_document_markdown(doc_id, markdown_content, "COMPLETED")
    
    # 2. 对 Markdown 内容进行切片
    chunks_text = parser_service.split_into_chunks(markdown_content)
    
    db_chunks = []
    qdrant_points = []
    
    # 3. 准备向量化与 upsert
    if chunks_text:
        # 批量向量化所有切片
        embeddings = embedding_service.get_embeddings_bulk(chunks_text, db=db)
        
        # 准备 Qdrant Collection
        collection_name = f"{settings.QDRANT_COLLECTION}_{kb_id}"
        qdrant_repository.init_collection(collection_name, len(embeddings[0]))
        
        for idx, (text, vector) in enumerate(zip(chunks_text, embeddings)):
            chunk_id = str(uuid.uuid4())
            # SQLite 实体
            chunk_model = ChunkModel(
                id=chunk_id,
                documentId=doc_id,
                knowledgeBaseId=kb_id,
                content=text,
                chunkIndex=idx,
                qdrantPointId=chunk_id  # 我们直接用 chunk_id 作为 Qdrant Point ID
            )
            db_chunks.append(chunk_model)
            
            # Qdrant Point
            qdrant_points.append(
                rest_models.PointStruct(
                    id=chunk_id,
                    vector=vector,
                    payload={
                        "documentId": doc_id,
                        "knowledgeBaseId": kb_id,
                        "chunkIndex": idx
                    }
                )
            )
            
        # 写入 SQLite Chunks
        repo.bulk_create_chunks(db_chunks)
        
        # 写入 Qdrant 向量
        qdrant_repository.upsert_vectors(collection_name, qdrant_points)
        
        # 4. 提取并存储 UnWeaver 实体关系
        for chunk in db_chunks:
            unweaver_service.extract_and_store_entities(db, kb_id, chunk.id, chunk.content)

@router.post("/upload", response_model=DocumentResponse, status_code=status.HTTP_201_CREATED)
def upload_document(
    kbId: str, 
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...), 
    db: Session = Depends(get_db)
):
    """
    上传文档，创建异步解析任务
    """
    repo = SQLiteRepository(db)
    
    # 校验知识库是否存在
    kb = repo.get_kb(kbId)
    if not kb:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Knowledge base not found"
        )
        
    # 保存上传的文件到本地
    doc_id = str(uuid.uuid4())
    file_extension = os.path.splitext(file.filename)[1]
    saved_file_name = f"{doc_id}{file_extension}"
    saved_file_path = os.path.join(UPLOAD_DIR, saved_file_name)
    
    with open(saved_file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    file_size = os.path.getsize(saved_file_path)
    
    # 创建数据库记录 (初始为 PENDING)
    doc_model = DocumentModel(
        id=doc_id,
        knowledgeBaseId=kbId,
        fileName=file.filename,
        filePath=saved_file_path,
        fileSize=file_size,
        status="PENDING"
    )
    doc = repo.create_document(doc_model)
    
    # 将解析和入库抛入后台异步执行，防止 HTTP 连接超时
    background_tasks.add_task(process_document_task, doc_id, kbId, saved_file_path)
    
    return doc

@router.get("/", response_model=list[DocumentResponse])
def list_documents(kbId: str, db: Session = Depends(get_db)):
    """
    获取知识库下的所有文档列表
    """
    repo = SQLiteRepository(db)
    return repo.get_documents_by_kb(kbId)

@router.get("/{doc_id}", response_model=DocumentResponse)
def get_document(doc_id: str, db: Session = Depends(get_db)):
    """
    获取单个文档信息
    """
    repo = SQLiteRepository(db)
    doc = repo.get_document(doc_id)
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found"
        )
    return doc

@router.get("/{doc_id}/markdown", response_model=dict)
def get_document_markdown(doc_id: str, db: Session = Depends(get_db)):
    """
    获取文档解析出的原始 Markdown 文本
    """
    repo = SQLiteRepository(db)
    doc = repo.get_document(doc_id)
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found"
        )
    return {"id": doc.id, "fileName": doc.fileName, "markdownContent": doc.markdownContent or ""}

@router.put("/{doc_id}/markdown", response_model=DocumentResponse)
def update_document_markdown(
    doc_id: str, 
    edit_req: DocumentEditRequest, 
    db: Session = Depends(get_db)
):
    """
    用户手工修改 Markdown 解析结果，保存后级联重新切片、生成向量和重新提取实体关系
    """
    repo = SQLiteRepository(db)
    doc = repo.get_document(doc_id)
    if not doc:
         raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found"
        )
         
    kb_id = doc.knowledgeBaseId
    collection_name = f"{settings.QDRANT_COLLECTION}_{kb_id}"
    
    # 1. 提取出原有的切片，获取对应的 Qdrant Point ID 列表，并进行清理
    old_chunks = repo.get_chunks_by_document(doc_id)
    old_point_ids = [c.id for c in old_chunks if c.id]
    
    # 删除旧的 Qdrant 向量
    if old_point_ids:
        try:
            qdrant_repository.delete_vectors(collection_name, old_point_ids)
        except Exception as e:
            print(f"Error clearing old vectors for document {doc_id}: {e}")
            
    # 2. 清理原有的数据库切片和实体关系数据
    # 为避免直接批量删除的逻辑错乱，我们通过 session 直接清理关联的模型
    db.query(ChunkModel).filter(ChunkModel.documentId == doc_id).delete()
    from app.model import EntityRelationModel
    db.query(EntityRelationModel).filter(EntityRelationModel.chunkId.in_(old_point_ids)).delete(synchronize_session=False)
    db.commit()
    
    # 3. 对编辑后的全新 Markdown 重新执行分词、向量化和关系挖掘入库
    try:
        index_document_markdown(db, repo, doc_id, kb_id, edit_req.markdownContent)
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to re-index edited markdown: {str(e)}"
        )
        
    return doc

@router.delete("/{doc_id}", status_code=status.HTTP_200_OK)
def delete_document(doc_id: str, db: Session = Depends(get_db)):
    """
    删除文档，并清空对应的 SQLite 切片、实体定义与 Qdrant 密集向量
    """
    repo = SQLiteRepository(db)
    doc = repo.get_document(doc_id)
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found"
        )
        
    kb_id = doc.knowledgeBaseId
    collection_name = f"{settings.QDRANT_COLLECTION}_{kb_id}"
    
    # 1. 清理 Qdrant 向量
    chunks = repo.get_chunks_by_document(doc_id)
    point_ids = [c.id for c in chunks]
    
    if point_ids:
        try:
            qdrant_repository.delete_vectors(collection_name, point_ids)
        except Exception as e:
            print(f"Error cleaning vectors for doc {doc_id}: {e}")
            
    # 2. 清理 UnWeaver 实体关系
    from app.model import EntityRelationModel
    try:
        db.query(EntityRelationModel).filter(EntityRelationModel.chunkId.in_(point_ids)).delete(synchronize_session=False)
    except Exception as e:
        print(f"Error cleaning entity relations for doc {doc_id}: {e}")
        
    # 3. 删除本地物理文件
    if doc.filePath and os.path.exists(doc.filePath):
        try:
            os.remove(doc.filePath)
        except Exception as e:
            print(f"Error deleting file {doc.filePath}: {e}")
            
    # 4. 删除 SQLite 中的文档记录 (级联删除 Chunks)
    success = repo.delete_document(doc_id)
    if not success:
         raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to delete document"
        )
         
    return {"message": "Document and associated vectors deleted successfully"}

@router.post("/parse-temp")
def parse_temporary_document(file: UploadFile = File(...)):
    """
    临时提取上传文档(PDF/TXT/MD/EXCEL/PPT)的纯文本，不存库，不产生向量，仅用于单次即时会话分析。
    """
    import tempfile
    
    filename = file.filename
    if not filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid filename"
        )
        
    ext = filename.split(".")[-1].lower()
    if ext not in ["pdf", "txt", "md", "markdown", "xlsx", "xls", "pptx", "ppt"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Supported: PDF, TXT, MD, EXCEL, PPT."
        )
        
    suffix = f".{ext}"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        try:
            content_bytes = file.file.read()
            tmp.write(content_bytes)
            tmp.close()  # 写入后关闭描述符以允许其他进程访问
            
            if ext == "pdf":
                text = parser_service.parse_pdf(tmp.name)
            elif ext in ["xlsx", "xls"]:
                import pandas as pd
                excel_file = pd.ExcelFile(tmp.name)
                sheets = []
                for sheet_name in excel_file.sheet_names:
                    df = pd.read_excel(excel_file, sheet_name=sheet_name)
                    df = df.fillna("")
                    md_table = df.to_markdown(index=False)
                    sheets.append(f"### 工作表 (Sheet): {sheet_name}\n\n{md_table}")
                text = "\n\n".join(sheets)
            elif ext in ["pptx", "ppt"]:
                from pptx import Presentation
                prs = Presentation(tmp.name)
                slides = []
                for idx, slide in enumerate(prs.slides):
                    slide_text = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_text.append(shape.text.strip())
                    slides.append(f"### 幻灯片第 {idx + 1} 页\n\n" + "\n".join(slide_text))
                text = "\n\n".join(slides)
            else:
                try:
                    text = content_bytes.decode("utf-8")
                except UnicodeDecodeError:
                    text = content_bytes.decode("gbk", errors="ignore")
                    
            return {"content": text, "fileName": filename}
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Temporary parsing failed: {str(e)}"
            )
        finally:
            if os.path.exists(tmp.name):
                os.remove(tmp.name)
