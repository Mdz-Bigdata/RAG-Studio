"""多模态 RAG：图片抽取 → 图片索引 → 文搜图 / 图搜图 / 图文混合 / 图搜文 → 带图片引用的回答。

全部离线：图片用 PIL 现场画，向量用离线感知特征；跨模态（CLIP）路径用一个"玩具双塔"验证逻辑。
"""

from __future__ import annotations

import base64
import io
import json
import os
from pathlib import Path

import numpy as np
import pytest
from PIL import Image, ImageDraw

from prism.config import Settings
from prism.generation import AnswerGenerator, FakeLLM
from prism.generation.answer import parse_image_citations
from prism.index import Index
from prism.ingest import IngestPipeline
from prism.multimodal import BaseImageEmbedder, ImageDescriber, MultimodalRetriever, PerceptualImageEmbedder, extract_images, wants_images
from prism.retrieval import HybridRetriever

REPORT = """# 2025 年 Q3 经营报告

## 销售额

第三季度各区域销售额如图所示，华东区 450 万，华北区 300 万，同比增长 15%。

![图1 Q3 各区域销售额柱状图](img/sales.png)

## 市场份额

公司产品的市场份额占比见下图，核心产品占 36%。

![图2 市场份额饼图](img/share.png)

## 系统架构

Ingress 网关把流量转发到后端服务，整体架构如下。

![图3 K8s Ingress 架构图](img/arch.png)

![装饰小图标](img/icon.png)
"""


def bar_chart() -> Image.Image:
    im = Image.new("RGB", (400, 300), "white")
    d = ImageDraw.Draw(im)
    for i, h in enumerate([120, 200, 160, 90]):
        d.rectangle([40 + i * 90, 280 - h, 100 + i * 90, 280], fill=(40, 90, 200))
    return im


def pie_chart() -> Image.Image:
    im = Image.new("RGB", (400, 300), "white")
    d = ImageDraw.Draw(im)
    d.pieslice([100, 40, 320, 260], 0, 130, fill=(220, 50, 50))
    d.pieslice([100, 40, 320, 260], 130, 360, fill=(250, 190, 40))
    return im


def arch_diagram() -> Image.Image:
    im = Image.new("RGB", (400, 300), (245, 245, 250))
    d = ImageDraw.Draw(im)
    for x, y in [(30, 30), (160, 30), (290, 30), (160, 180)]:
        d.rectangle([x, y, x + 90, y + 60], outline=(0, 0, 0), width=3)
    d.line([75, 90, 205, 180], fill=(0, 0, 0), width=2)
    d.line([335, 90, 205, 180], fill=(0, 0, 0), width=2)
    return im


def data_url(im: Image.Image, fmt: str = "PNG") -> str:
    buf = io.BytesIO()
    im.save(buf, format=fmt)
    return f"data:image/{fmt.lower()};base64," + base64.b64encode(buf.getvalue()).decode()


@pytest.fixture(scope="module")
def corpus(tmp_path_factory) -> Path:
    root = tmp_path_factory.mktemp("mm_corpus")
    (root / "img").mkdir()
    bar_chart().save(root / "img" / "sales.png")
    pie_chart().save(root / "img" / "share.png")
    arch_diagram().save(root / "img" / "arch.png")
    Image.new("RGB", (16, 16), "red").save(root / "img" / "icon.png")
    (root / "report.md").write_text(REPORT, encoding="utf-8")
    return root


@pytest.fixture(scope="module")
def mm_index(corpus: Path, tmp_path_factory) -> Index:
    settings = Settings.lite(index_dir=tmp_path_factory.mktemp("mm_index"))
    pipeline = IngestPipeline(settings)
    report = pipeline.ingest_paths([corpus / "report.md"])
    assert not report.failed, report.failed
    return pipeline.index


def captions(hits) -> list[str]:
    return [h.record.caption for h in hits]


# --------------------------------------------------------------------------------------
def test_extract_and_link(mm_index: Index):
    recs = mm_index.images.of_document(next(iter(mm_index.store.docs)))
    assert [r.caption for r in recs] == ["图1 Q3 各区域销售额柱状图", "图2 市场份额饼图", "图3 K8s Ingress 架构图"]  # 16px 图标被过滤
    assert [r.breadcrumbs[-1] for r in recs] == ["销售额", "市场份额", "系统架构"]
    assert all(r.anchor_chunk_id and mm_index.store.get(r.anchor_chunk_id).kind == "parent" for r in recs)
    assert all(mm_index.images.path_of(r).is_file() for r in recs)
    # 上下文不跨小节：饼图的摘要里不应出现销售额小节的"华东区"
    assert "华东区" in recs[0].search_text() and "华东区" not in recs[1].search_text()
    assert mm_index.stats()["images"] == 3


def test_text_to_image(mm_index: Index):
    r = MultimodalRetriever(mm_index.settings, mm_index)
    res = r.text_to_image("各区域销售额柱状图")
    assert captions(res.images)[0].startswith("图1")
    assert set(res.routes) == {"desc_dense", "desc_sparse"}  # 离线：没有文本塔，CLIP 路跳过并说明
    assert any("没有文本塔" in n for n in res.notes)
    top, second = res.images[0], res.images[1]
    assert top.score == pytest.approx(1.0) and second.score < 0.8  # 相对相关度能拉开差距
    assert captions(r.text_to_image("Ingress 网关 架构").images)[0].startswith("图3")


def test_image_to_image(mm_index: Index):
    r = MultimodalRetriever(mm_index.settings, mm_index)
    query = pie_chart().resize((200, 150))  # 同一张图的缩小版
    res = r.image_to_image(query)
    assert captions(res.images)[0].startswith("图2") and res.images[0].score > 0.95
    assert res.images[1].score < 0.8
    assert len(r.image_to_image(query, min_score=0.9).images) == 1  # 相似度阈值过滤


def test_hybrid_late_fusion(mm_index: Index):
    r = MultimodalRetriever(mm_index.settings, mm_index)
    pie = pie_chart()
    assert captions(r.hybrid("架构图", pie, alpha=0.9).images)[0].startswith("图3")   # 文本主导
    assert captions(r.hybrid("架构图", pie, alpha=0.1).images)[0].startswith("图2")   # 图片主导


def test_image_to_text(mm_index: Index):
    r = MultimodalRetriever(mm_index.settings, mm_index)
    res = r.image_to_text(bar_chart())
    assert res.contexts and "华东区 450 万" in res.contexts[0].chunk.text   # 相似图片 → 所在段落
    assert "neighbor:image→image→chunk" in res.routes
    assert any("所在的段落排在最前" in n for n in res.notes) and res.fused[0].chunk_id in mm_index.store.get(
        mm_index.images.of_document(next(iter(mm_index.store.docs)))[0].anchor_chunk_id).child_ids
    # 附带了问题：问题可能指向别处，不强行置顶
    assert not any("所在的段落排在最前" in n for n in r.image_to_text(bar_chart(), query="市场份额").notes)
    # 没有任何信号（陌生图片 + 无 OCR / VLM / 文本塔）→ 空结果并说明缺什么
    res = r.image_to_text(Image.effect_noise((200, 200), 80).convert("RGB"))
    assert not res.contexts and any("没有可用信号" in n for n in res.notes)


def test_images_for_answer_gating(mm_index: Index):
    hr = HybridRetriever(mm_index.settings, mm_index)
    mr = MultimodalRetriever(mm_index.settings, mm_index)

    def pick(q: str) -> list[str]:
        hits, _ = mr.images_for_answer(q, hr.retrieve(q, top_k=3).contexts)
        return captions(hits)

    assert [c[:2] for c in pick("华东区销售额是多少")] == ["图1"]        # 图文关联 + 关键词证据
    assert pick("年假可以跨年使用吗") == []                              # 不相关的问题不挂图
    assert wants_images("给我看一下架构图") and not wants_images("年假可以跨年吗")
    assert [c[:2] for c in pick("给我看一下架构图")] == ["图3"]          # 显式要图


def test_persistence_and_reembed(mm_index: Index):
    mm_index.save()
    loaded = Index.load(mm_index.settings, mm_index.directory)
    assert len(loaded.images) == 3 and len(loaded.images.visual) == 3
    res = MultimodalRetriever(loaded.settings, loaded).image_to_image(arch_diagram())
    assert captions(res.images)[0].startswith("图3")

    # 换了图文向量模型：加载后用磁盘上的原图自动重算
    toy = ToyCrossModal()
    swapped = Index.load(mm_index.settings, mm_index.directory, image_embedder=toy)
    res = MultimodalRetriever(swapped.settings, swapped).text_to_image("饼图")
    assert swapped.images.visual.dim == toy.dim and "clip" in res.routes


def test_relink_keeps_images_and_remove(corpus: Path, tmp_path: Path):
    settings = Settings.lite(index_dir=tmp_path / "idx")
    p = IngestPipeline(settings)
    doc = p.parse(corpus / "report.md")
    p.ingest_document(doc, source_path=corpus / "report.md")
    assert len(p.index.images) == 3
    # 校对回灌：不带源文件重新切片 → 图片保留并重新挂接
    edited = p.parse_markdown_text(REPORT.replace("## 市场份额", "## 市场占有率"), doc_id=doc.doc_id, source=str(corpus / "report.md"))
    p.ingest_document(edited)
    recs = p.index.images.of_document(doc.doc_id)
    assert len(recs) == 3 and recs[1].breadcrumbs[-1] == "市场占有率"
    p.index.remove_document(doc.doc_id)
    assert len(p.index.images) == 0 and not (p.index.directory / "images" / doc.doc_id).exists()


# --------------------------------------------------------------------------------------
def test_extract_pdf_docx_pptx_and_image_file(tmp_path: Path):
    cfg = Settings.lite().multimodal
    bar, pie = bar_chart(), pie_chart()

    pdf = tmp_path / "charts.pdf"
    bar.save(pdf, format="PDF", save_all=True, append_images=[pie])
    got = extract_images(pdf, None, cfg)
    assert [g.page for g in got] == [0, 1] and got[0].origin == "pdf"

    import docx
    from docx.shared import Inches

    bar.save(tmp_path / "bar.png")
    d = docx.Document()
    d.add_paragraph("第三季度销售情况如下图。")
    d.add_picture(str(tmp_path / "bar.png"), width=Inches(3))
    d.add_paragraph("图 1 各区域销售额")
    d.save(tmp_path / "r.docx")
    got = extract_images(tmp_path / "r.docx", None, cfg)
    assert len(got) == 1 and got[0].caption == "图 1 各区域销售额" and "第三季度" in got[0].context

    from pptx import Presentation
    from pptx.util import Inches as PInches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "市场份额"
    pie.save(tmp_path / "pie.png")
    slide.shapes.add_picture(str(tmp_path / "pie.png"), PInches(1), PInches(1.5), width=PInches(4))
    prs.save(tmp_path / "deck.pptx")
    got = extract_images(tmp_path / "deck.pptx", None, cfg)
    assert len(got) == 1 and got[0].page == 0 and "市场份额" in (got[0].caption + got[0].context)

    # 图片文件本身就是一篇文档：正文 = 图片文本摘要，图片挂在这个块上
    settings = Settings.lite(index_dir=tmp_path / "idx")
    p = IngestPipeline(settings)
    arch_diagram().save(tmp_path / "Ingress架构图.png")
    report = p.ingest_paths([tmp_path / "Ingress架构图.png", pdf], save=False)
    assert not report.failed and report.parsers.get("image") == 1
    rec = next(r for r in p.index.images.records.values() if r.origin == "file")
    assert rec.anchor_chunk_id and "Ingress架构图" in p.index.store.get(rec.anchor_chunk_id).text
    assert len(p.index.images) == 3


def test_repeated_logo_is_dropped(tmp_path: Path):
    logo, chart = Image.new("RGB", (120, 120), (10, 120, 60)), bar_chart()
    pdf = tmp_path / "logo.pdf"
    # 每页一张图：logo 出现在 3 页，图表只出现 1 页
    logo.save(pdf, format="PDF", save_all=True, append_images=[logo, logo, chart])
    got = extract_images(pdf, None, Settings.lite().multimodal)
    assert len(got) == 1 and got[0].image.size == (400, 300)


# --------------------------------------------------------------------------------------
class ToyCrossModal(BaseImageEmbedder):
    """玩具双塔：图片走感知特征；文本按关键词映射到对应参考图的向量——模拟"图文对齐"的空间。"""

    name = "toy-clip"
    model_name = "toy"
    cross_modal = True

    def __init__(self) -> None:
        self.inner = PerceptualImageEmbedder()
        self.dim = self.inner.dim
        self.refs = {"销售": bar_chart(), "柱状": bar_chart(), "份额": pie_chart(), "饼图": pie_chart(), "架构": arch_diagram(), "Ingress": arch_diagram()}

    def embed_images(self, images):
        return self.inner.embed_images(images)

    def embed_texts(self, texts):
        out = []
        for t in texts:
            hit = [self.inner.embed_image(im) for k, im in self.refs.items() if k in t]
            v = np.mean(hit, axis=0) if hit else np.full(self.dim, 1e-3, dtype=np.float32)
            out.append(v / (np.linalg.norm(v) or 1.0))
        return np.asarray(out, dtype=np.float32)


@pytest.fixture(scope="module")
def clip_index(corpus: Path, tmp_path_factory) -> Index:
    settings = Settings.lite(index_dir=tmp_path_factory.mktemp("clip_index"))
    index = Index(settings, image_embedder=ToyCrossModal())
    IngestPipeline(settings, index=index).ingest_paths([corpus / "report.md"], save=False)
    return index


def test_cross_modal_routes(clip_index: Index):
    r = MultimodalRetriever(clip_index.settings, clip_index)
    res = r.text_to_image("市场份额饼图")
    assert set(res.routes) == {"clip", "desc_dense", "desc_sparse"} and captions(res.images)[0].startswith("图2")
    assert res.images[0].routes["clip"].rank == 1

    # 早融合：q = α·text + (1-α)·image
    bar = bar_chart()
    res = r.hybrid("架构", bar, alpha=0.95)
    assert "clip_fused" in res.routes and captions(res.images)[0].startswith("图3")
    assert res.routes["clip_fused"][0].chunk_id == res.images[0].record.image_id
    assert r.hybrid("架构", bar, alpha=0.0).routes["clip_fused"][0].chunk_id.endswith("img000")

    # 图搜文直达路：图片向量 → 正文块的图文空间向量（入库时已随模型增量建好）
    assert len(clip_index.images.chunk_clip) == len(clip_index.store.child_ids())
    res = r.image_to_text(pie_chart())
    assert "clip:image→chunk" in res.routes and "市场份额" in res.contexts[0].chunk.text


def test_describer_ocr_vlm_cache(tmp_path: Path):
    cfg = Settings.lite().multimodal.model_copy(update={"vlm_caption": True})
    vlm = FakeLLM(lambda p: "一张展示 Q3 各区域销售额的蓝色柱状图", vision=True)
    d = ImageDescriber(cfg, vlm, cache_path=tmp_path / "cache.json")
    assert d.capabilities == {"ocr": None, "vlm": "fake"}
    ocr, desc = d.describe(bar_chart(), "sha-a")
    assert ocr == "" and "柱状图" in desc and vlm.image_counts == [1]
    d2 = ImageDescriber(cfg, vlm, cache_path=tmp_path / "cache.json")
    assert d2.caption(bar_chart(), "sha-a") == desc and len(vlm.calls) == 1   # 命中缓存，不再调用
    assert ImageDescriber(cfg, FakeLLM()).vlm is None                         # 看不了图的模型不会被用来描述图片


def test_vlm_description_is_indexed(corpus: Path, tmp_path: Path):
    settings = Settings.lite(index_dir=tmp_path / "idx")
    settings.multimodal.vlm_caption = True
    vlm = FakeLLM(lambda p: "折线走势示意：营收环比上升", vision=True)
    p = IngestPipeline(settings, vlm=vlm)
    p.ingest_paths([corpus / "report.md"], save=False)
    assert all("折线走势" in r.description for r in p.index.images.records.values())
    assert MultimodalRetriever(settings, p.index).text_to_image("营收环比").images[0].routes["desc_sparse"].rank <= 3


def test_answer_with_image_references(mm_index: Index):
    hr = HybridRetriever(mm_index.settings, mm_index)
    mr = MultimodalRetriever(mm_index.settings, mm_index)
    q = "华东区销售额是多少，请出示相关图表"
    result = hr.retrieve(q, top_k=3)
    hits, _ = mr.images_for_answer(q, result.contexts)
    assert hits and hits[0].record.caption.startswith("图1")

    # 离线抽取式回答也会列出图片
    offline = AnswerGenerator(mm_index.settings.generation, None).answer(result, images=hits)
    assert "[图1]" in offline.text

    llm = FakeLLM(lambda p: "根据 [图1] 所示，华东区销售额为 450 万 [1]。", vision=True)
    gen = AnswerGenerator(mm_index.settings.generation, llm)
    ans = gen.answer(result, images=hits, image_payloads=[data_url(bar_chart())], image_b64=data_url(pie_chart()))
    prompt = llm.calls[-1]
    assert "<图片资料>" in prompt and "[图1] 来源：" in prompt and "图注：图1 Q3 各区域销售额柱状图" in prompt
    assert "<附图顺序>用户图片、[图1]</附图顺序>" in prompt
    assert llm.image_counts[-1] == 2 and parse_image_citations(ans.text, len(hits)) == [1] and ans.citations == [1]

    # 文本模型看不了图：不发像素，改用自动识别结果说明用户图片
    text_llm = FakeLLM(lambda p: "ok")
    AnswerGenerator(mm_index.settings.generation, text_llm).answer(result, images=hits, image_payloads=[data_url(bar_chart())], image_b64=data_url(pie_chart()), query_image_note="饼图，图中文字：36%")
    assert text_llm.image_counts[-1] == 0 and "<用户图片>" in text_llm.calls[-1] and "36%" in text_llm.calls[-1]


# --------------------------------------------------------------------------------------
@pytest.fixture(scope="module")
def client(tmp_path_factory):
    data_dir = tmp_path_factory.mktemp("mm_appdata")
    os.environ["SQLITE_URL"] = f"sqlite:///{data_dir}/app.db"
    os.environ["OPENAI_API_KEY"] = ""
    os.environ["ANTHROPIC_API_KEY"] = ""
    from fastapi.testclient import TestClient

    from app import config as app_config_mod
    from app.services import index_service, model_service, settings_service

    app_config_mod.app_config.data_dir = data_dir
    app_config_mod.app_config.upload_dir = data_dir / "uploads"
    app_config_mod.app_config.index_root = data_dir / "indexes"
    for d in (app_config_mod.app_config.upload_dir, app_config_mod.app_config.index_root):
        d.mkdir(parents=True, exist_ok=True)
    for svc in (settings_service, model_service, index_service):
        svc.invalidate()

    from app.main import app

    with TestClient(app) as c:
        # 测试环境固定用离线感知特征（本机即使装了 torch 也不去加载 / 下载 CLIP）
        r = c.put("/api/v1/settings/", json={"overrides": {"multimodal": {"embedding": {"provider": "perceptual"}, "ocr": "none"}, "parsing": {"mineru": {"enabled": False}}}})
        assert r.status_code == 200, r.text
        # 保存配置会在后台线程检查"哪些知识库要重建"。等它结束再开始上传：否则它可能在刚上传完时把新库判成
        # "索引不完整"、排队重建，文档又回到处理中，"全部 COMPLETED"的断言偶发失败（整套测试里出现过一次）
        import time

        deadline = time.time() + 60
        while c.get("/api/v1/settings/reindex-status").json()["busy"] and time.time() < deadline:
            time.sleep(0.1)
        yield c
        c.post("/api/v1/settings/reset")


def sse_events(raw: str) -> list[tuple[str, object]]:
    out = []
    for block in raw.split("\n\n"):
        if block.strip():
            name, data = block.split("\ndata: ", 1)
            out.append((name.replace("event: ", ""), json.loads(data)))
    return out


def test_api_multimodal_end_to_end(client):
    kb_id = client.post("/api/v1/knowledge-bases/", json={"name": "经营报告", "category": "报告"}).json()["id"]
    # Markdown 内嵌 data URI 图片（Web 上传时本地相对路径不可用）+ 直接上传图片文件
    md = REPORT.replace("img/sales.png", data_url(bar_chart())).replace("img/share.png", data_url(pie_chart())).replace("(img/arch.png)", "(img/missing.png)").replace("![装饰小图标](img/icon.png)", "")
    assert client.post(f"/api/v1/documents/upload?kbId={kb_id}", files={"file": ("report.md", md.encode(), "text/markdown")}).status_code == 201
    buf = io.BytesIO()
    arch_diagram().save(buf, format="PNG")
    assert client.post(f"/api/v1/documents/upload?kbId={kb_id}", files={"file": ("K8s Ingress 架构图.png", buf.getvalue(), "image/png")}).status_code == 201
    docs = client.get(f"/api/v1/documents/?kbId={kb_id}").json()
    assert all(d["status"] == "COMPLETED" for d in docs), [(d["fileName"], d["errorMessage"]) for d in docs]
    # 内嵌图片的 base64 不进正文 / 切片
    report = next(d for d in docs if d["fileName"] == "report.md")
    stored = client.get(f"/api/v1/documents/{report['id']}/markdown").json()["markdownContent"]
    assert "base64" not in stored and "![图1 Q3 各区域销售额柱状图](embedded-image)" in stored
    assert max(c["tokenCount"] for c in client.get(f"/api/v1/documents/{report['id']}/chunks").json()) < 200

    st = client.get(f"/api/v1/multimodal/{kb_id}/status").json()
    assert st["images"] == 3 and st["documentsWithImages"] == 2 and st["embedder"]["provider"] == "perceptual" and st["embedder"]["crossModal"] is False
    gallery = client.get(f"/api/v1/multimodal/{kb_id}/images").json()
    assert len(gallery) == 3 and {g["fileName"] for g in gallery} == {"report.md", "K8s Ingress 架构图.png"}
    f = client.get(gallery[0]["url"])
    assert f.status_code == 200 and f.headers["content-type"].startswith("image/")
    assert client.get(f"/api/v1/multimodal/{kb_id}/images/nope/file").status_code == 404

    search = lambda body: client.post(f"/api/v1/multimodal/{kb_id}/search", json=body)  # noqa: E731
    # 文搜图
    r = search({"mode": "text2image", "query": "各区域销售额柱状图"}).json()
    assert r["images"][0]["caption"].startswith("图1") and set(r["routes"]) == {"desc_dense", "desc_sparse"} and r["routes"]["desc_sparse"][0]["kind"] == "image"
    assert search({"mode": "text2image", "query": "Ingress 架构图"}).json()["images"][0]["fileName"] == "K8s Ingress 架构图.png"
    # 图搜图
    r = search({"mode": "image2image", "imageB64": data_url(pie_chart().resize((180, 135)), "JPEG"), "topK": 2}).json()
    assert r["images"][0]["caption"].startswith("图2") and r["images"][0]["score"] > 0.9 and len(r["images"]) == 2
    # 图文混合
    r = search({"mode": "hybrid", "query": "架构图", "imageB64": data_url(pie_chart()), "alpha": 0.9}).json()
    assert r["images"][0]["fileName"] == "K8s Ingress 架构图.png" and "visual" in r["routes"]
    # 图搜文
    r = search({"mode": "image2text", "imageB64": data_url(bar_chart())}).json()
    assert r["contexts"] and "华东区 450 万" in r["contexts"][0]["text"] and r["images"][0]["caption"].startswith("图1")
    # 参数校验
    assert search({"mode": "image2image"}).status_code == 400
    assert search({"mode": "hybrid", "query": "x"}).status_code == 400
    assert search({"mode": "image2image", "imageB64": "not-an-image"}).status_code == 400

    def chat(body):
        with client.stream("POST", "/api/v1/chat/stream", json={"kbId": kb_id, **body}) as resp:
            assert resp.status_code == 200
            return sse_events("".join(resp.iter_text()))

    # 端到端：文字提问 → 回答带图片引用
    ev = chat({"message": "华东区销售额是多少，请出示相关图表", "topK": 3})
    names = [n for n, _ in ev]
    assert names[:5] == ["meta", "plan", "sources", "relations", "images"] and names[-1] == "done"
    images = dict(ev)["images"]
    assert images and images[0]["index"] == 1 and images[0]["caption"].startswith("图1") and images[0]["linked"] is True
    text = "".join(d["content"] for n, d in ev if n == "message")
    assert "[图1]" in text and ev[-1][1]["imageCitations"] == [1]
    assert dict(chat({"message": "给我看一下 Ingress 架构图", "topK": 2}))["images"][0]["fileName"] == "K8s Ingress 架构图.png"
    # 不相关的问题不挂图
    assert dict(chat({"message": "年假可以跨年使用吗", "topK": 2}))["images"] == []

    # 只发图片：以图搜图 + 以图搜文
    ev = chat({"message": "", "imageB64": data_url(pie_chart())})
    meta, plan = ev[0][1], ev[1][1]
    assert meta["queryImage"] and client.get(meta["queryImage"]).status_code == 200
    assert plan["strategy"] == "image" and plan["imageSearch"]["mode"] == "image2text+image2image"
    assert dict(ev)["images"][0]["caption"].startswith("图2")
    assert any("市场份额" in s["text"] for s in dict(ev)["sources"])
    assert ev[-1][1]["vision"] is False
    # 图片 + 文字：图文混合
    ev = chat({"message": "和这张图相关的销售数据", "imageB64": data_url(bar_chart())})
    assert ev[1][1]["strategy"].endswith("+image") and ev[1][1]["imageSearch"]["mode"] == "image2text+hybrid"
    assert any("华东区 450 万" in s["text"] for s in dict(ev)["sources"])

    # 历史消息里能回看查询图片与回答图片
    msgs = client.get(f"/api/v1/chat/sessions/{meta['sessionId']}/messages").json()
    assert msgs[0]["content"] == "（图片检索）" and msgs[0]["trace"]["queryImage"] == meta["queryImage"]
    assert msgs[1]["trace"]["images"][0]["caption"].startswith("图2")
    assert client.post("/api/v1/chat/stream", json={"kbId": kb_id, "message": ""}).status_code == 400

    # 运行状态与模型配置
    rt = client.get("/api/v1/system/status").json()["runtime"]
    assert rt["imageEmbedding"]["provider"] == "perceptual" and rt["vision"] is False and rt["multimodalEnabled"] is True
    assert any("image_embedding" in p["categories"] for p in client.get("/api/v1/models/providers").json())
    cfg = client.post("/api/v1/models/configs", json={"providerName": "Jina 多模态向量 API", "providerType": "api", "baseUrl": "https://api.jina.ai/v1", "apiKey": "jina_test_1234567890", "category": "image_embedding", "modelName": "jina-clip-v2"}).json()
    assert cfg["category"] == "image_embedding"
    assert client.delete(f"/api/v1/models/configs/{cfg['id']}").status_code == 200

    # 人工校对过的文档：校对回灌保留图片；整库重建（旧索引先挪走）后从原始文件重新抽取，图片不丢
    from app.services import document_service

    edited_md = client.get(f"/api/v1/documents/{report['id']}/markdown").json()["markdownContent"].replace("## 市场份额", "## 市场占有率")
    assert client.put(f"/api/v1/documents/{report['id']}/markdown", json={"markdownContent": edited_md}).status_code == 200
    assert client.get(f"/api/v1/multimodal/{kb_id}/status").json()["images"] == 3
    assert document_service.reindex_kb(kb_id) == []
    gallery_after = client.get(f"/api/v1/multimodal/{kb_id}/images").json()
    assert len(gallery_after) == 3 and any(g["breadcrumb"].endswith("市场占有率") for g in gallery_after)
    assert client.get(f"/api/v1/documents/{report['id']}").json()["parser"].endswith("+edited")
    gallery = gallery_after

    # 多模态评测：自动样例 → 运行 → 历史
    assert len(client.get(f"/api/v1/eval/{kb_id}/multimodal/sample").json()) >= 8
    rep = client.post(f"/api/v1/eval/{kb_id}/multimodal/run", json={"auto": True, "ks": [1, 3]}).json()
    assert set(rep["task_metrics"]) == {"text2image", "image2image", "image2text", "answer"} and rep["markdown"].startswith("# 多模态评测")
    assert client.get(f"/api/v1/eval/{kb_id}/multimodal/reports").json()[0]["createdAt"] == rep["createdAt"]

    # 删除文档 → 图片与文件一起清理
    doc = next(d for d in docs if d["fileName"] == "report.md")
    assert client.delete(f"/api/v1/documents/{doc['id']}").status_code == 200
    assert client.get(f"/api/v1/multimodal/{kb_id}/status").json()["images"] == 1
    assert client.get(gallery[0]["url"]).status_code == 404 or gallery[0]["fileName"] != "report.md"
    assert client.delete(f"/api/v1/knowledge-bases/{kb_id}").status_code == 200


# --------------------------------------------------------------------------------------
# 带图表的真实形态 PDF：页眉 logo 每页重复、图注在图下方、最后一页是没有文字层的扫描件
def _report_pdf(path: Path) -> None:
    pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    w, h = A4
    logo = Image.new("RGB", (220, 80), (20, 60, 140))
    scan = Image.new("RGB", (900, 1200), (246, 244, 238))
    c = canvas.Canvas(str(path), pagesize=A4)
    pages = [("一、区域销售情况 华东区继续领跑，见图1。", bar_chart(), "图1 2025年Q3各区域销售额"), ("二、产品结构 收入构成见图2。", pie_chart(), "图2 产品收入构成")]
    for body, chart, caption in pages:
        c.drawImage(ImageReader(logo), 40, h - 60, 70, 25)
        c.setFont("STSong-Light", 12)
        c.drawString(60, h - 100, body)
        c.drawImage(ImageReader(chart), 80, h - 420, 400, 300)
        c.drawString(80, h - 440, caption)
        c.showPage()
    c.drawImage(ImageReader(logo), 40, h - 60, 70, 25)
    c.drawImage(ImageReader(scan), 40, 60, w - 80, h - 140)
    c.showPage()
    c.save()


def test_report_pdf_captions_logo_and_scanned_page(tmp_path: Path, monkeypatch):
    pdf = tmp_path / "report.pdf"
    _report_pdf(pdf)
    settings = Settings.lite(index_dir=tmp_path / "idx")
    p = IngestPipeline(settings)
    # 扫描页的 OCR：本测试不依赖本机 OCR，按图片尺寸伪造识别结果
    monkeypatch.setattr(p.describer.ocr, "read", lambda im: "巡检单号 XJ-2025-0931 设备编号 PUMP-07 轴承温度偏高" if im.size[1] > im.size[0] else "")
    doc = p.parse(pdf)
    p.ingest_document(doc, source_path=pdf)
    recs = p.index.images.of_document(doc.doc_id)
    # logo 每页都有 → 丢弃，而且不占"图1"的号：图注按顺序落在真正的图表上
    assert [(r.page, r.caption) for r in recs] == [(0, "图1 2025年Q3各区域销售额"), (1, "图2 产品收入构成"), (2, "")]
    # 扫描页没有文字层：OCR 结果回填为正文，文本检索也能找到
    assert any(b.meta.get("ocr_backfill") and b.page == 2 for b in doc.blocks)
    hit = HybridRetriever(settings, p.index).retrieve("PUMP-07 出了什么问题", top_k=1, use_transform=False).contexts[0]
    assert "PUMP-07" in hit.chunk.text

    mr = MultimodalRetriever(settings, p.index)
    hr = HybridRetriever(settings, p.index)

    def attached(q):
        return [h.record.page for h in mr.images_for_answer(q, hr.retrieve(q, top_k=2, use_transform=False).contexts)[0]]

    assert attached("华东区销售额是多少") == [0]       # 只带有关键词证据的那张
    assert attached("PUMP-07 出了什么问题") == [2]
    assert attached("年假可以跨年使用吗") == []         # "年假"被切成"年 / 假"也不会虚高


def test_ocr_layout_keeps_chart_labels_with_values():
    from prism.multimodal.describe import layout_ocr

    # rapidocr 对柱状图的真实框（数值标签随柱高上下错开）：按阅读顺序拼会把华东区配成 380
    boxes = [(333, 32, 821, 71, "2025年Q3各区域销售额"), (806, 28, 967, 78, "（万元)"), (436, 112, 520, 155, "450"), (677, 188, 758, 233, "380"),
             (198, 277, 278, 318, "300"), (916, 377, 1000, 420, "210"), (190, 673, 301, 714, "华北区"), (429, 670, 542, 718, "华东区"),
             (669, 670, 782, 718, "华南区"), (910, 674, 1022, 713, "西部区")]
    assert layout_ocr(boxes, 1200) == "2025年Q3各区域销售额 （万元) 300 华北区 | 450 华东区 | 380 华南区 | 210 西部区"
    # 没有列结构（扫描件 / 截图）：保持阅读顺序
    lines = [(0, 0, 900, 40, "巡检单号：XJ-2025-0931"), (0, 60, 700, 100, "设备编号：PUMP-07"), (0, 120, 300, 160, "巡检人：王工"), (400, 120, 700, 160, "审核：李经理")]
    assert layout_ocr(lines, 1000) == "巡检单号：XJ-2025-0931 设备编号：PUMP-07 巡检人：王工 审核：李经理"


def test_multimodal_eval_runner(mm_index: Index):
    from prism.eval import MultimodalEvalRunner, MultimodalExample, build_samples
    from prism.eval.multimodal import match_image

    recs = mm_index.images.of_document(next(iter(mm_index.store.docs)))
    assert match_image(recs[1], "市场份额饼图") and match_image(recs[1], recs[1].image_id) and not match_image(recs[0], "饼图")

    samples = build_samples(mm_index.images, mm_index.store)
    tasks = {s.task for s in samples}
    assert tasks == {"text2image", "image2image", "image2text", "answer"} and sum(s.expect_no_image for s in samples) == 3
    # 图搜文的关键词来自各图自己的图注，而不是三张图都一样
    assert len({tuple(s.relevant_keywords) for s in samples if s.task == "image2text"}) == 3

    runner = MultimodalEvalRunner(MultimodalRetriever(mm_index.settings, mm_index), HybridRetriever(mm_index.settings, mm_index), ks=[1, 3])
    report = runner.run(samples + [
        MultimodalExample(task="text2image", question="K8s Ingress 架构", relevant_images=["Ingress 架构图"]),
        MultimodalExample(task="text2image", question="完全无关的问题", relevant_images=["不存在的图"]),
        MultimodalExample(task="image2image", query_image="kb:不存在", relevant_images=["x"]),
    ])
    m = report.task_metrics
    # 3 张同图扰动全中；3 张模拟翻拍里离线感知特征把翻拍的饼图认成了架构图（桌面背景和光照压过了颜色 / 线条统计，
    # 换成 CLIP 后同样的翻拍全部排第一）；坏样例记为未命中。翻拍样例的意义正在于暴露这类差距
    assert m["image2image"]["hit@1"] == pytest.approx(5 / 7)
    assert {e.category for e in report.examples if e.task == "image2image"} >= {"同图扰动", "模拟翻拍"}
    assert m["text2image"]["hit@3"] < 1.0 and m["answer"]["负例不附图率"] == 1.0
    assert set(report.route_metrics) == {"desc_dense", "desc_sparse", "fused"}
    bad = [e for e in report.examples if not e.ok]
    assert any("未匹配到任何图片" in e.note for e in bad) and any("出错" in e.note and "不存在" in e.note for e in bad)
    md = report.to_dict()["markdown"]
    assert "## 文搜图" in md and "## 回答附图" in md and "未通过的样例" in md


def test_exif_orientation_is_applied():
    from prism.multimodal.extract import open_image

    # 横着存的像素 + EXIF Orientation=6（顺时针转 90° 才是正的）：手机竖拍照片的典型形态
    raw = Image.new("RGB", (400, 200), "white")
    ImageDraw.Draw(raw).rectangle([0, 0, 50, 200], fill="red")
    exif = Image.Exif()
    exif[0x0112] = 6
    buf = io.BytesIO()
    raw.save(buf, format="JPEG", exif=exif)
    im = open_image(buf.getvalue())
    assert im.size == (200, 400)                       # 已摆正成竖图
    assert im.getpixel((100, 10))[0] > 200 and im.getpixel((100, 10))[1] < 80   # 红条转到了顶部


def test_chunk_vectors_are_built_in_batches(clip_index: Index):
    images = clip_index.images
    children = [clip_index.store.chunks[cid] for cid in clip_index.store.child_ids()]
    assert len(children) >= 3
    images.chunk_clip.remove(list(images.chunk_clip.ids))          # 模拟大库：一个向量都还没建
    added, remaining = images.sync_chunk_vectors(children, limit=2)
    assert (added, remaining) == (2, len(children) - 2)
    images.chunk_clip.remove(list(images.chunk_clip.ids))
    r = MultimodalRetriever(clip_index.settings.with_overrides({"multimodal": {"chunk_sync_per_query": 1}}), clip_index)
    res = r.image_to_text(pie_chart())                                 # 每次查询只补 1 条，并如实说明
    assert any("还在补建" in n for n in res.notes) and len(images.chunk_clip) == 1
    added, remaining = images.sync_chunk_vectors(children)             # 不限量：一次补齐
    assert remaining == 0 and len(images.chunk_clip) == len(children)


def test_scanned_page_backfill_needs_a_full_page_image():
    from prism.ingest.pipeline import IngestPipeline as P
    from prism.multimodal.extract import ExtractedImage
    from prism.schemas import Block, BlockType, DocMetadata, ParsedDocument

    def doc():
        return ParsedDocument(doc_id="d", metadata=DocMetadata(title="报告"), blocks=[
            Block(type=BlockType.TITLE, text="一、区域销售情况", page=0),
            Block(type=BlockType.TEXT, text="第三季度公司实现销售额 1340 万元，同比增长 15%，各区域表现分化明显，具体见图1。", page=0),
        ])

    ocr = "2025年Q3各区域销售额 300 华北区 | 450 华东区 | 380 华南区 | 210 西部区"
    big = Image.new("RGB", (1200, 800))
    # MinerU 给出 bbox：图表只占页面一小块 → 不是扫描页（实测误判过的形态）
    d = doc()
    P._backfill_scanned_pages(d, [ExtractedImage(image=big, sha1="a", origin="mineru", page=0, bbox=[203, 286, 822, 552], ocr_text=ocr)])
    assert not any(b.meta.get("ocr_backfill") for b in d.blocks)
    # 整页大图 → 扫描页，回填
    d = doc()
    P._backfill_scanned_pages(d, [ExtractedImage(image=big, sha1="a", origin="mineru", page=0, bbox=[30, 40, 970, 960], ocr_text=ocr)])
    assert any(b.meta.get("ocr_backfill") for b in d.blocks)
    # 没有位置信息：页面有正文就不回填；几乎空白才回填
    d = doc()
    P._backfill_scanned_pages(d, [ExtractedImage(image=big, sha1="a", origin="pdf", page=0, ocr_text=ocr)])
    assert not any(b.meta.get("ocr_backfill") for b in d.blocks)
    empty = ParsedDocument(doc_id="e", metadata=DocMetadata(), blocks=[Block(type=BlockType.TEXT, text="ACME 集团", page=0)])
    P._backfill_scanned_pages(empty, [ExtractedImage(image=big, sha1="a", origin="pdf", page=0, ocr_text=ocr)])
    assert any(b.meta.get("ocr_backfill") for b in empty.blocks)



def test_image_to_text_accepts_clear_winner_below_threshold(clip_index: Index):
    """翻拍 / 截图后的图和原图相似度过不了 0.9 的门槛，但明显领先第二名时仍认定是同一张图；
    第一名不够高、或领先不够多（同风格的无关图）时不认定。"""
    r = MultimodalRetriever(clip_index.settings, clip_index)
    emb = clip_index.images.embedder
    target = clip_index.images.of_document(next(iter(clip_index.store.docs)))[1]   # 饼图
    vec = clip_index.images.visual.get_vectors([target.image_id])[target.image_id]

    def run(scores):
        ids = list(clip_index.images.visual.ids)
        ranked = [(target.image_id, scores[0])] + [(i, s) for i, s in zip([x for x in ids if x != target.image_id], scores[1:])]
        orig = clip_index.images.visual.search
        clip_index.images.visual.search = lambda *a, **k: ranked
        try:
            return r.image_to_text(pie_chart(), describe=False)
        finally:
            clip_index.images.visual.search = orig

    assert emb.similar_threshold == 0.9
    res = run([0.78, 0.45, 0.40])                         # 实测翻拍形态：领先 0.33
    assert "neighbor:image→image→chunk" in res.routes and any("明显领先" in n for n in res.notes)
    assert "neighbor:image→image→chunk" not in run([0.66, 0.59, 0.50]).routes   # 同风格无关图：领先不够
    assert "neighbor:image→image→chunk" not in run([0.53, 0.35, 0.25]).routes   # 第一名太低


def test_edited_pdf_rebuild_keeps_captions(tmp_path: Path):
    """校对过的 PDF 在整库重建时：正文用校对稿，图片从原文件重新抽取，并且图注 / 页码按原文件结构对上。"""
    pdf = tmp_path / "report.pdf"
    _report_pdf(pdf)
    settings = Settings.lite(index_dir=tmp_path / "idx")
    p = IngestPipeline(settings)
    original = p.parse(pdf)
    edited = p.parse_markdown_text(original.to_markdown() + "\n\n## 校对补充\n\n人工补充的一段说明。", doc_id=original.doc_id, source=str(pdf))
    # 不传 image_doc：Markdown 没有页码，图注对不上
    p.ingest_document(edited, source_path=pdf)
    assert all(not r.caption for r in p.index.images.of_document(original.doc_id))
    # 传原文件结构：图注回来了，正文仍是校对稿
    p.index.remove_document(original.doc_id)
    p.ingest_document(p.parse_markdown_text(edited.markdown, doc_id=original.doc_id, source=str(pdf)), source_path=pdf, image_doc=p.parse(pdf))
    recs = p.index.images.of_document(original.doc_id)
    assert [(r.page, r.caption) for r in recs][:2] == [(0, "图1 2025年Q3各区域销售额"), (1, "图2 产品收入构成")]
    assert any("人工补充的一段说明" in c.text for c in p.index.store.chunks.values())


def test_number_code_pattern():
    from prism.multimodal.retriever import _NUM_CODE

    found = lambda q: _NUM_CODE.findall(q)  # noqa: E731
    assert found("华中防火墙表格截图3-11") == ["3-11"] and found("工单卡片50-20") == ["50-20"]
    assert found("版本 1.2.3 发布") == ["1.2.3"] and found("日期 2025/09/18") == ["2025/09/18"]
    assert found("单价 3.5 元") == [] and found("2025年销售") == []
    assert found("巡检单 XJ-2025-0931") == []            # 字母开头的编号走分词匹配，不重复算


def test_figure_label_lookup(mm_index: Index):
    """点名图号 / 表号（"Figure 20""表 2""图3"）：图注以该编号开头的图排最前，回答时也算作强证据。"""
    from prism.multimodal.retriever import figure_labels

    assert figure_labels("Figure 20 和 Fig. 3 说明了什么") == {("figure", 20), ("figure", 3)}
    assert figure_labels("表2 的数据") == {("table", 2)} and figure_labels("看一下图 3") == {("figure", 3)}
    assert figure_labels("图表3-11") == set() and figure_labels("2025年图表") == set()     # 数字分段码 / "图表"不是图号
    assert figure_labels("Figure 3. Numpy-like pseudocode", caption=True) == {("figure", 3)}
    assert figure_labels("Accuracy compared with Figure 5", caption=True) == set()       # 图注正文里的引用不算本图编号
    r = MultimodalRetriever(mm_index.settings, mm_index)
    res = r.text_to_image("图3")
    assert captions(res.images)[0].startswith("图3") and any("图号" in n for n in res.notes)
    hits, _ = r.images_for_answer("图2 说明了什么", [])
    assert captions(hits) == ["图2 市场份额饼图"]


def test_sparse_route_weight_follows_query_coverage(mm_index: Index):
    """查询的词在图片文本里几乎都没有（中文提问 + 英文图注）：关键词路降权，也不再拿它的第一名去"互相印证"。"""
    r = MultimodalRetriever(mm_index.settings, mm_index)
    assert r._sparse_coverage("各区域销售额") > 0.9
    assert r._sparse_coverage("pseudocode contrastive loss") == 0.0
    _, weights = r._desc_routes("市场份额 pseudocode contrastive loss objective", None)
    assert 0 < weights["desc_sparse"] < mm_index.settings.multimodal.weight_desc_sparse
    _, weights = r._desc_routes("市场份额饼图", None)
    assert weights["desc_sparse"] == mm_index.settings.multimodal.weight_desc_sparse


def test_clip_query_baseline_needs_enough_images(clip_index: Index):
    """问题一侧的 CLIP 基线（各图超出中性基线的中位数）：图片太少时中位数就是候选本身，不启用。"""
    r = MultimodalRetriever(clip_index.settings, clip_index)
    res = r.text_to_image("市场份额")
    assert "clip" in res.routes and len(clip_index.images) < 8
    assert r._clip_query_baseline("市场份额", res) == 0.0


def test_vlm_description_goes_before_parser_transcript(corpus: Path, tmp_path: Path, monkeypatch):
    """MinerU flash 档给的"图片内容"是图中文字转写：VLM 的语义描述放前面，转写接在后面，两者都保留。"""
    from prism.ingest import pipeline as pipeline_mod
    from prism.multimodal.extract import ExtractedImage, pixel_sha1

    im = bar_chart()
    fake = lambda *a, **k: [ExtractedImage(image=im, sha1=pixel_sha1(im), origin="mineru", page=0, caption="Figure 1. Summary", description="(1) Contrastive pre-training")]  # noqa: E731
    monkeypatch.setattr(pipeline_mod, "extract_images", fake)
    settings = Settings.lite(index_dir=tmp_path / "idx")
    settings.multimodal.vlm_caption = True
    p = IngestPipeline(settings, vlm=FakeLLM(lambda q: "这是一张对比预训练方法的流程示意图", vision=True))
    [it] = p.extract_images(corpus / "report.md", p.parse(corpus / "report.md"))
    assert it.description.startswith("这是一张对比预训练方法的流程示意图") and it.description.endswith("(1) Contrastive pre-training")
    p_no_vlm = IngestPipeline(Settings.lite(index_dir=tmp_path / "idx2"))
    [it] = p_no_vlm.extract_images(corpus / "report.md", p_no_vlm.parse(corpus / "report.md"))
    assert it.description == "(1) Contrastive pre-training"                         # 没有 VLM 时保持解析器的文字


def test_ocr_respace_restores_english_word_gaps():
    """rapidocr 的中文模型不输出英文空格：按逐字框中心距补回；中文、数字之间、紧挨着的字符不受影响。"""
    from prism.multimodal.describe import OCREngine, _respace

    def boxes(centers):
        return [[[c - 3, 0], [c + 3, 0], [c + 3, 10], [c - 3, 10]] for c in centers]

    # 字距 8px，"the" 前跳了 15px → 补空格
    assert _respace(list("Pepperthe"), boxes([0, 8, 16, 24, 32, 40, 55, 63, 71])) == "Pepper the"
    assert _respace(list("best(model)"), boxes([0, 8, 16, 24, 40, 48, 56, 64, 72, 80, 88])) == "best (model)"
    # 宽字母把中心距撑大不算词间隔：M→a、F→e、m→a 的中心距比别处大 1/8~3/8，按字宽归一化后不拆；e→F 跳了 18px 才是词间隔
    assert _respace(list("MaleFemale"), boxes([0, 11, 19, 27, 45, 54, 62, 71, 79, 87])) == "Male Female"
    assert _respace(list("ValueA1234567"), boxes([0, 8, 16, 24, 32, 40, 60, 68, 76, 84, 92, 100, 108])) is None  # 拉丁字母不到 8 个：不处理
    assert _respace(list("54.7%"), boxes([0, 8, 30, 38, 46])) is None                    # 数字 / 中文行不处理
    assert _respace(list("华东区450"), boxes([0, 20, 40, 70, 78, 86])) is None
    assert _respace(list("Pepper the"), boxes(range(0, 80, 8))) is None                  # 模型已经输出了空格：信模型
    assert _respace([], []) is None and _respace(list("abc"), boxes([0, 8])) is None

    font_path = "/System/Library/Fonts/Helvetica.ttc"
    if OCREngine("rapidocr").name != "rapidocr" or not os.path.exists(font_path):
        pytest.skip("没有 rapidocr 或英文字体")
    from PIL import ImageFont

    im = Image.new("RGB", (450, 80), "white")
    ImageDraw.Draw(im).text((20, 25), "Contrastive pre-training of image encoders", font=ImageFont.truetype(font_path, 18), fill="black")
    text = OCREngine("rapidocr").read(im)
    # 这一行模型自己输出了部分空格：信模型、不再补（标准答案测试里这样整体更准），也不会把单词从中间拆开
    assert "of image encoders" in text and "Contrastive" in text and "im age" not in text, text


def test_cross_lingual_query_translation(mm_index: Index):
    """问题和图片文字语言不同、关键词路匹配不上时，用辅助模型把问题译成图片文字的语言补充检索；同语言不调用模型。"""
    from prism.multimodal import retriever as retriever_mod

    retriever_mod._TRANSLATIONS.clear()
    llm = FakeLLM(lambda p: "各区域销售额柱状图")
    r = MultimodalRetriever(mm_index.settings, mm_index, llm=llm)
    res = r.text_to_image("regional sales bar chart")
    assert captions(res.images)[0].startswith("图1") and any("译为简体中文" in n for n in res.notes)
    hits, _ = r.images_for_answer("what were the regional sales", [])
    assert captions(hits)[:1] == ["图1 Q3 各区域销售额柱状图"]           # 译文的词作为附图的关键词证据
    calls = len(llm.calls)
    r.text_to_image("各区域销售额")                                     # 同语言：不译
    r.text_to_image("regional sales bar chart")                        # 已译过：走缓存
    assert len(llm.calls) == calls
    echo = MultimodalRetriever(mm_index.settings, mm_index, llm=FakeLLM())   # 原样回显的"译文"不当作补充
    assert echo._translate("market share pie") == ""
    off = mm_index.settings.model_copy(deep=True)
    off.multimodal.cross_lingual = False
    assert MultimodalRetriever(off, mm_index, llm=llm)._translate("regional sales bar chart") == ""
    retriever_mod._TRANSLATIONS.clear()


def test_mineru_figure_fragments_inherit_caption_and_inner_label():
    """复合图被拆成几块时，同页连续、只有一块带图注的碎片继承图注；图注前混进图表标题时仍认得出本图编号。"""
    from prism.multimodal.extract import _from_mineru
    from prism.multimodal.retriever import figure_labels
    from prism.schemas import Block, BlockType, DocMetadata, ParsedDocument

    img = lambda im, cap="", page=0: Block(type=BlockType.IMAGE, page=page, meta={"img_path": data_url(im), "caption": cap})  # noqa: E731
    doc = ParsedDocument(doc_id="d", metadata=DocMetadata(source="/tmp/x.pdf"), blocks=[
        Block(type=BlockType.TEXT, text="正文一段。", page=0),
        img(bar_chart()), img(pie_chart()), img(arch_diagram(), "Robustness Figure 14. While supervised adaptation"),
        Block(type=BlockType.TEXT, text="下一段正文。", page=0),
        img(bar_chart().rotate(90, expand=True)),                       # 被正文隔开：另一组，不继承
        img(pie_chart().rotate(90, expand=True), page=1),               # 换页：另一组
    ])
    caps = [x.caption for x in _from_mineru(doc, None)]
    assert caps[:3] == ["Robustness Figure 14. While supervised adaptation"] * 3 and caps[3:] == ["", ""]
    assert figure_labels(caps[0], caption=True) == {("figure", 14)}
    assert figure_labels("Accuracy compared with Figure 5", caption=True) == set()
