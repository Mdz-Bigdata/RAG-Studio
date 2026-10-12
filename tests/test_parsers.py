from __future__ import annotations

from datetime import date
from pathlib import Path

from prism.ingest.parsers.excel_parser import ExcelParser
from prism.ingest.parsers.html_table import html_table_to_grid, html_table_to_markdown
from prism.ingest.parsers.markdown_parser import MarkdownParser
from prism.ingest.parsers.mineru_parser import content_list_to_blocks
from prism.schemas import BlockType

RAW = Path(__file__).resolve().parents[1] / "data" / "raw"


def test_markdown_front_matter_and_structure():
    doc = MarkdownParser().parse(RAW / "员工手册_2026版.md")
    assert doc.metadata.title == "员工手册（2026 版）"
    assert doc.metadata.version == "2026.1"
    assert doc.metadata.valid_from == date(2026, 1, 1)
    assert "制度" in doc.metadata.tags
    titles = [(b.level, b.text) for b in doc.blocks if b.type == BlockType.TITLE]
    assert (1, "员工手册（2026 版）") in titles
    assert (3, "年假") in titles
    tables = [b for b in doc.blocks if b.type == BlockType.TABLE]
    assert tables and tables[0].text.startswith("| 职级")
    lists = [b for b in doc.blocks if b.type == BlockType.LIST]
    assert any("出差结束后 10 个工作日内" in b.text for b in lists)


def test_markdown_code_fence(tmp_path: Path):
    p = tmp_path / "x.md"
    p.write_text("# T\n\n```python\nprint(1)\n```\n\n段落。\n", encoding="utf-8")
    doc = MarkdownParser().parse(p)
    kinds = [b.type for b in doc.blocks]
    assert kinds == [BlockType.TITLE, BlockType.CODE, BlockType.TEXT]
    assert doc.blocks[1].meta["language"] == "python"


def test_excel_csv_row_blocks():
    doc = ExcelParser().parse(RAW / "产品价格表.csv")
    rows = [b for b in doc.blocks if b.type == BlockType.TEXT]
    assert len(rows) == 5
    assert rows[2].meta["record"]["芯片"] == "M4 Pro"
    assert "售价（元）: 15999" in rows[2].text
    assert any(b.type == BlockType.TABLE for b in doc.blocks)


def test_html_table_rowspan_colspan():
    html = "<table><tr><th>A</th><th colspan='2'>B</th></tr><tr><td rowspan='2'>x</td><td>1</td><td>2</td></tr><tr><td>3</td><td>4</td></tr></table>"
    grid = html_table_to_grid(html)
    assert grid == [["A", "B", "B"], ["x", "1", "2"], ["x", "3", "4"]]
    md = html_table_to_markdown(html)
    assert md.splitlines()[0] == "| A | B | B |"


def test_content_list_to_blocks_mapping():
    items = [
        {"type": "text", "text": "第一章 总则", "text_level": 1, "page_idx": 0, "bbox": [10, 10, 500, 40]},
        {"type": "text", "text": "正文段落。", "page_idx": 0, "bbox": [10, 50, 500, 80]},
        {"type": "table", "table_caption": ["表 1"], "table_body": "<table><tr><td>a</td><td>b</td></tr></table>", "table_footnote": [], "img_path": "images/t.png", "page_idx": 1, "bbox": [0, 0, 1000, 1000]},
        {"type": "page_number", "text": "3", "page_idx": 1, "bbox": [0, 0, 1, 1]},
        {"type": "list", "list_items": ["甲", "乙"], "page_idx": 1, "bbox": [0, 0, 1, 1]},
    ]
    blocks = content_list_to_blocks(items)
    assert [b.type for b in blocks] == [BlockType.TITLE, BlockType.TEXT, BlockType.TABLE, BlockType.PAGE_NUMBER, BlockType.LIST]
    assert blocks[0].level == 1 and blocks[0].page == 0
    assert blocks[2].text.startswith("表 1\n| a | b |")


def test_mineru_tier_overview(monkeypatch):
    from prism.config import MinerUConfig
    from prism.ingest.parsers import mineru_tiers as mt

    # 本机未安装：四档仍全部列出且可选
    ov = mt.tier_overview(MinerUConfig(tier="basic"), available=False)
    assert [t["tier"] for t in ov["tiers"]] == ["flash", "basic", "standard", "advanced"]
    assert all(t["selectable"] and t["status"] == "not_installed" for t in ov["tiers"]) and ov["source"] == "catalog"
    assert [t["tier"] for t in ov["tiers"] if t["current"]] == ["basic"]

    # 已安装：flash 直接可用，其余逐档核对模型包（advanced 与 standard 共用）
    probed = []
    monkeypatch.setattr(mt, "local_tier_status", lambda tier: (probed.append(tier), {"status": "ready" if tier in ("flash", "basic") else "needs_models", "repos": [{"name": "m", "ready": tier == "basic", "root": "/x", "missing": 0}], "missingModules": []})[1])
    ov = mt.tier_overview(MinerUConfig(), available=True)
    assert [t["status"] for t in ov["tiers"]] == ["ready", "ready", "needs_models", "needs_models"] and probed == ["flash", "basic", "standard", "advanced"]
    assert ov["tiers"][1]["modelRepos"][0]["ready"] is True

    # 自托管服务：以 /v1/tiers 为准，并接受目录之外的新档位
    assert mt._parse_remote_tiers({"tiers": [{"tier": "flash", "available": True}, {"name": "standard", "ready": False}, "turbo"]}) == {"flash": True, "standard": False, "turbo": True}
    monkeypatch.setattr(mt, "remote_tiers", lambda url, key="", timeout=3.0: {"url": url, "reachable": True, "tiers": {"flash": True, "standard": True, "turbo": True}, "error": None})
    ov = mt.tier_overview(MinerUConfig(api_url="http://mineru:8000"), available=False)
    status = {t["tier"]: t["status"] for t in ov["tiers"]}
    assert status == {"flash": "remote", "basic": "remote_missing", "standard": "remote", "advanced": "remote_missing", "turbo": "remote"} and ov["source"] == "remote"

    # 硬件建议
    assert mt.recommend_tier({"accelerator": "apple_silicon"})["tier"] == "standard"
    assert mt.recommend_tier({"accelerator": "nvidia", "gpuMemoryGb": 6})["tier"] == "basic"
    assert mt.recommend_tier({"accelerator": "nvidia", "gpuMemoryGb": 24})["extra"] == "full"
    assert mt.recommend_tier({"accelerator": "none"})["tier"] == "basic"


def test_mineru_failure_falls_back_to_builtin_parser(tmp_path, monkeypatch):
    import docx

    from prism.config import Settings
    from prism.ingest import IngestPipeline
    from prism.ingest.parsers.base import BaseParser

    class Broken(BaseParser):
        name = "mineru"

        def parse(self, path, metadata=None):
            raise RuntimeError("model package incomplete")

    d = docx.Document()
    d.add_heading("报销制度", 1)
    d.add_paragraph("差旅住宿标准为每晚 500 元。")
    d.save(tmp_path / "a.docx")
    p = IngestPipeline(Settings.lite(index_dir=tmp_path / "idx"))
    monkeypatch.setattr(p.registry, "get_parser", lambda path: Broken())
    doc = p.parse(tmp_path / "a.docx")
    assert doc.metadata.parser == "docx" and "model package incomplete" in doc.metadata.extra["mineru_error"]
    assert any("500 元" in b.text for b in doc.blocks)


def test_mineru_standard_double_bullet_is_collapsed():
    from prism.ingest.parsers.mineru_parser import content_list_to_blocks

    # MinerU 4.0.10 standard 档对真实课程大纲 PDF 的输出形状
    items = [{"type": "index", "list_items": ["- 课程大纲", "- - 第1集:架构设计与技术选型 ★", "  - - 系统架构图设计"], "page_idx": 0}]
    assert content_list_to_blocks(items)[0].text.split("\n") == ["- 课程大纲", "- 第1集:架构设计与技术选型 ★", "  - 系统架构图设计"]


def test_mineru_chart_caption_in_footnote():
    from prism.ingest.parsers.mineru_parser import content_list_to_blocks

    # MinerU 4.0.10 standard 档对图表 PDF 的实际输出：图注落在 chart_footnote 里
    b = content_list_to_blocks([{"type": "chart", "content": "", "chart_caption": [], "chart_footnote": ["图1 2025年Q3各区域销售额"], "img_path": "data:image/png;base64,AAAA", "page_idx": 0}])[0]
    assert b.meta["caption"] == "图1 2025年Q3各区域销售额" and b.text == "图1 2025年Q3各区域销售额"
