"""PPTX 兜底解析器（无 MinerU 时使用）：python-pptx，每页幻灯片一个标题块 + 文本 / 表格块。"""

from __future__ import annotations

from pathlib import Path

from ...schemas import Block, BlockType, DocMetadata, ParsedDocument
from .base import BaseParser
from .html_table import grid_to_markdown


class PptxParser(BaseParser):
    name = "pptx"
    extensions = frozenset({"pptx"})

    def parse(self, path: Path, metadata: DocMetadata | None = None) -> ParsedDocument:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("解析 pptx 需要 python-pptx（或安装 MinerU）") from exc
        doc = self._new_document(path, metadata)
        prs = Presentation(str(path))
        for idx, slide in enumerate(prs.slides):
            title = ""
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
            doc.blocks.append(Block(type=BlockType.TITLE, level=2, text=f"幻灯片 {idx + 1}" + (f"：{title}" if title else ""), page=idx))
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if getattr(shape, "has_table", False) and shape.has_table:
                    grid = [[c.text.strip().replace("\n", " ") for c in row.cells] for row in shape.table.rows]
                    doc.blocks.append(Block(type=BlockType.TABLE, text=grid_to_markdown(grid), page=idx, meta={"format": "markdown"}))
                elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                    if text:
                        doc.blocks.append(Block(type=BlockType.TEXT, text=text, page=idx))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    doc.blocks.append(Block(type=BlockType.TEXT, text=f"演讲备注：{notes}", page=idx))
        return doc
