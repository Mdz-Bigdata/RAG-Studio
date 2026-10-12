"""Docx 兜底解析器（无 MinerU 时使用）：python-docx，按 Heading 样式恢复层级，表格转 Markdown。"""

from __future__ import annotations

import re
from pathlib import Path

from ...schemas import Block, BlockType, DocMetadata, ParsedDocument
from .base import BaseParser
from .html_table import grid_to_markdown

_HEADING_STYLE = re.compile(r"heading\s*(\d)", re.I)


class DocxParser(BaseParser):
    name = "docx"
    extensions = frozenset({"docx"})

    def parse(self, path: Path, metadata: DocMetadata | None = None) -> ParsedDocument:
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("解析 docx 需要 python-docx（或安装 MinerU）：pip install 'rag-agent-studio[office]'") from exc
        doc = self._new_document(path, metadata)
        d = docx.Document(str(path))
        body = d.element.body
        para_map = {p._p: p for p in d.paragraphs}
        table_map = {t._tbl: t for t in d.tables}
        for child in body.iterchildren():
            if child in para_map:
                p = para_map[child]
                text = p.text.strip()
                if not text:
                    continue
                style = (p.style.name if p.style is not None else "") or ""
                if style.lower() == "title":
                    doc.blocks.append(Block(type=BlockType.TITLE, level=1, text=text))
                    continue
                m = _HEADING_STYLE.search(style)
                if m:
                    doc.blocks.append(Block(type=BlockType.TITLE, level=int(m.group(1)), text=text))
                elif "list" in style.lower():
                    doc.blocks.append(Block(type=BlockType.LIST, text="- " + text))
                else:
                    doc.blocks.append(Block(type=BlockType.TEXT, text=text))
            elif child in table_map:
                t = table_map[child]
                grid = [[c.text.strip().replace("\n", " ") for c in row.cells] for row in t.rows]
                doc.blocks.append(Block(type=BlockType.TABLE, text=grid_to_markdown(grid), meta={"format": "markdown"}))
        return doc
