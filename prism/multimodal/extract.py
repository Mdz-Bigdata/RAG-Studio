"""从文档里抽取图片，并记录它在文档中的位置（页码 / 邻近正文 / 图注）。

位置信息必须在抽取时就留下——生成回答时才知道"这是第几页的哪张图"，图文关联也靠它。

支持：
- PDF        PyMuPDF（若已安装）→ 否则 pypdf 的 page.images
- DOCX       段落里的内嵌图片（a:blip），上下文取前后段落
- PPTX       每页幻灯片的图片形状（含组合形状），上下文取幻灯片标题与文本
- Markdown   ![alt](path "title") 与 <img>：本地相对路径 / data URI（不抓取 http 图片）
- MinerU     Content List 里的 image / chart / table 块（img_path 指向已落盘的图片）
- 图片文件   png / jpg / jpeg / webp / bmp / gif / tif 本身就是一张图

过滤：太小的图（图标 / 分隔线）、同一张图在多页重复出现（logo / 水印）、文档内重复。
"""

from __future__ import annotations

import base64
import hashlib
import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ..config import MultimodalConfig
from ..schemas import BlockType, ParsedDocument

logger = logging.getLogger(__name__)

IMAGE_EXTENSIONS = frozenset({"png", "jpg", "jpeg", "webp", "bmp", "gif", "tif", "tiff"})

_MD_IMAGE = re.compile(r"!\[([^\]]*)\]\(\s*<?([^)\s>]+)>?(?:\s+[\"']([^\"']*)[\"'])?\s*\)")
_MD_HEADING = re.compile(r"^#{1,6}\s", re.M)
_HTML_IMAGE = re.compile(r"<img\b[^>]*>", re.I)
_HTML_ATTR = re.compile(r"(\w+)\s*=\s*[\"']([^\"']*)[\"']")
_CAPTION = re.compile(r"^\s*(?:图|圖|图表|Figure|Fig\.?|Chart)\s*[\dA-Za-z一二三四五六七八九十]+[\s.．:：、\-—]*\S", re.I)


@dataclass
class ExtractedImage:
    image: Any                           # PIL.Image（RGB）
    sha1: str
    origin: str
    page: int | None = None
    bbox: list[float] | None = None
    caption: str = ""
    context: str = ""
    markup: str = ""                     # Markdown 里的原始图片标记，用于精确定位所在块
    ocr_text: str = ""
    description: str = ""
    meta: dict[str, Any] = field(default_factory=dict)


# --------------------------------------------------------------------------------------
def open_image(data: bytes) -> Any | None:
    """bytes → PIL.Image（RGB，透明通道铺白底，按 EXIF 方向摆正）。无法解码返回 None。"""
    from PIL import Image, ImageOps

    try:
        im = Image.open(io.BytesIO(data))
        im.load()
        # 手机照片的像素是横着存的，靠 EXIF 标记"该转 90°"：不摆正的话向量和 OCR 看到的都是躺倒的图
        im = ImageOps.exif_transpose(im) or im
    except Exception:
        return None
    return to_rgb(im)


def to_rgb(im: Any) -> Any:
    from PIL import Image

    if getattr(im, "is_animated", False):
        im.seek(0)
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        rgba = im.convert("RGBA")
        bg = Image.new("RGB", rgba.size, (255, 255, 255))
        bg.paste(rgba, mask=rgba.split()[-1])
        return bg
    return im.convert("RGB")


def decode_image_b64(image_b64: str) -> Any:
    """`data:image/png;base64,...` 或裸 base64 → PIL.Image。"""
    raw = image_b64.split(",", 1)[1] if image_b64.strip().startswith("data:") and "," in image_b64 else image_b64
    try:
        data = base64.b64decode(raw.strip().replace("\n", "") + "=" * (-len(raw.strip()) % 4))
    except Exception as exc:
        raise ValueError("图片 base64 无法解码") from exc
    im = open_image(data)
    if im is None:
        raise ValueError("无法识别的图片格式")
    return im


def pixel_sha1(im: Any) -> str:
    return hashlib.sha1(im.tobytes() + str(im.size).encode()).hexdigest()


def _clip(text: str, n: int) -> str:
    text = " ".join((text or "").split())
    return text[:n]


# --------------------------------------------------------------------------------------
def extract_images(path: Path, doc: ParsedDocument | None, cfg: MultimodalConfig, *, mineru_out_dir: Path | None = None) -> list[ExtractedImage]:
    """按文件类型抽取图片并过滤。任何一步失败都只记日志——图片抽取不应让文档入库失败。"""
    ext = path.suffix.lower().lstrip(".")
    try:
        items: list[ExtractedImage] = []
        if doc is not None and doc.metadata.parser.startswith("mineru"):
            items = _from_mineru(doc, mineru_out_dir)
        if not items:
            if ext in IMAGE_EXTENSIONS:
                items = _from_image_file(path)
            elif ext == "pdf":
                items = _from_pdf(path, doc, cfg)
            elif ext == "docx":
                items = _from_docx(path, cfg)
            elif ext == "pptx":
                items = _from_pptx(path, cfg)
            elif ext in ("md", "markdown"):
                items = extract_markdown_images(path.read_text(encoding="utf-8", errors="replace"), path.parent, cfg)
    except Exception as exc:
        logger.warning("图片抽取失败 %s：%s", path.name, exc)
        return []
    return filter_images(items, cfg)


def filter_images(items: list[ExtractedImage], cfg: MultimodalConfig) -> list[ExtractedImage]:
    pages_of: dict[str, set[int | None]] = {}
    for it in items:
        pages_of.setdefault(it.sha1, set()).add(it.page)
    out: list[ExtractedImage] = []
    seen: set[str] = set()
    for it in items:
        w, h = it.image.size
        if min(w, h) < cfg.min_image_side or w * h < cfg.min_image_area:
            continue
        if it.sha1 in seen:
            continue
        if cfg.drop_repeated_pages > 1 and len(pages_of[it.sha1] - {None}) >= cfg.drop_repeated_pages:
            continue  # 每页都出现的 logo / 水印
        seen.add(it.sha1)
        if max(w, h) > cfg.max_side:
            it.image = it.image.copy()
            it.image.thumbnail((cfg.max_side, cfg.max_side))
        out.append(it)
        if len(out) >= cfg.max_images_per_doc:
            logger.info("图片数超过上限 %d，其余忽略", cfg.max_images_per_doc)
            break
    return out


# --------------------------------------------------------------------------------------
def _from_image_file(path: Path) -> list[ExtractedImage]:
    im = open_image(path.read_bytes())
    if im is None:
        return []
    return [ExtractedImage(image=im, sha1=pixel_sha1(im), origin="file", caption=path.stem)]


def _page_texts(doc: ParsedDocument | None) -> dict[int, list[str]]:
    pages: dict[int, list[str]] = {}
    if doc is None:
        return pages
    for b in doc.blocks:
        if b.page is not None and b.text.strip() and b.type not in (BlockType.HEADER, BlockType.FOOTER, BlockType.PAGE_NUMBER):
            pages.setdefault(b.page, []).append(b.text.strip())
    return pages


def _page_captions(texts: list[str]) -> list[str]:
    caps: list[str] = []
    for t in texts:
        for line in t.splitlines():
            if _CAPTION.match(line) and len(line.strip()) <= 120:
                caps.append(line.strip())
    return caps


def _from_pdf(path: Path, doc: ParsedDocument | None, cfg: MultimodalConfig) -> list[ExtractedImage]:
    raw: list[tuple[int, bytes, list[float] | None]] = []
    try:
        import fitz  # PyMuPDF：能拿到图片在页面上的位置

        with fitz.open(str(path)) as pdf:
            for page_idx, page in enumerate(pdf):
                for info in page.get_images(full=True):
                    xref = info[0]
                    try:
                        data = pdf.extract_image(xref)["image"]
                    except Exception:
                        continue
                    rects = page.get_image_rects(xref)
                    bbox = None
                    if rects and page.rect.width and page.rect.height:
                        r = rects[0]
                        bbox = [r.x0 / page.rect.width, r.y0 / page.rect.height, r.x1 / page.rect.width, r.y1 / page.rect.height]
                    raw.append((page_idx, data, bbox))
    except ImportError:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        for page_idx, page in enumerate(reader.pages):
            try:
                images = list(page.images)
            except Exception as exc:  # 个别滤镜 pypdf 解不了
                logger.debug("第 %d 页图片无法解码：%s", page_idx + 1, exc)
                continue
            for img in images:
                raw.append((page_idx, img.data, None))

    pages = _page_texts(doc)
    decoded = [(page_idx, im, bbox) for page_idx, data, bbox in raw if (im := open_image(data)) is not None]
    # 图注按顺序分配给"会被保留"的图：页眉 logo、小图标不占图注的号（否则 logo 抢走"图1"，后面全部错位）
    sha_pages: dict[str, set[int]] = {}
    for page_idx, im, _ in decoded:
        sha_pages.setdefault(pixel_sha1(im), set()).add(page_idx)
    used: dict[int, int] = {}
    out: list[ExtractedImage] = []
    for page_idx, im, bbox in decoded:
        sha = pixel_sha1(im)
        texts = pages.get(page_idx, [])
        decorative = min(im.size) < cfg.min_image_side or im.size[0] * im.size[1] < cfg.min_image_area or (cfg.drop_repeated_pages > 1 and len(sha_pages[sha]) >= cfg.drop_repeated_pages)
        caption = ""
        if not decorative:
            caps = _page_captions(texts)
            k = used.get(page_idx, 0)
            used[page_idx] = k + 1
            caption = caps[k] if k < len(caps) else ""
        out.append(ExtractedImage(image=im, sha1=sha, origin="pdf", page=page_idx, bbox=bbox, caption=caption, context=_clip("\n".join(texts), cfg.context_chars)))
    return out


def _from_docx(path: Path, cfg: MultimodalConfig) -> list[ExtractedImage]:
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(str(path))
    paras = d.paragraphs
    texts = [p.text.strip() for p in paras]
    out: list[ExtractedImage] = []
    for i, p in enumerate(paras):
        blips = p._p.findall(".//" + qn("a:blip"))
        for blip in blips:
            rid = blip.get(qn("r:embed"))
            part = d.part.related_parts.get(rid) if rid else None
            im = open_image(part.blob) if part is not None else None
            if im is None:
                continue
            before = next((t for t in reversed(texts[:i]) if t), "")
            after = next((t for t in texts[i + 1:] if t), "")
            caption = texts[i] if _CAPTION.match(texts[i]) else after if _CAPTION.match(after) else before if _CAPTION.match(before) else ""
            out.append(ExtractedImage(image=im, sha1=pixel_sha1(im), origin="docx", caption=caption, context=_clip("\n".join(x for x in (before, texts[i], after) if x), cfg.context_chars)))
    return out


def _from_pptx(path: Path, cfg: MultimodalConfig) -> list[ExtractedImage]:
    from pptx import Presentation

    def walk(shapes):
        for sh in shapes:
            if getattr(sh, "shapes", None) is not None and sh.shape_type == 6:  # GROUP
                yield from walk(sh.shapes)
            else:
                yield sh

    prs = Presentation(str(path))
    out: list[ExtractedImage] = []
    for idx, slide in enumerate(prs.slides):
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        texts = [sh.text_frame.text.strip() for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip()]
        for sh in walk(slide.shapes):
            try:
                blob = sh.image.blob
            except Exception:
                continue
            im = open_image(blob)
            if im is None:
                continue
            descr = ""
            try:
                descr = sh._element.xpath(".//p:cNvPr/@descr")[0]
            except Exception:
                pass
            out.append(ExtractedImage(image=im, sha1=pixel_sha1(im), origin="pptx", page=idx, caption=descr or title, context=_clip("\n".join(texts), cfg.context_chars)))
    return out


def extract_markdown_images(text: str, base_dir: Path | None, cfg: MultimodalConfig) -> list[ExtractedImage]:
    found: list[tuple[int, int, str, str, str]] = []  # (start, end, src, alt, markup)
    for m in _MD_IMAGE.finditer(text):
        found.append((m.start(), m.end(), m.group(2), m.group(1) or m.group(3) or "", m.group(0)))
    for m in _HTML_IMAGE.finditer(text):
        attrs = {k.lower(): v for k, v in _HTML_ATTR.findall(m.group(0))}
        if attrs.get("src"):
            found.append((m.start(), m.end(), attrs["src"], attrs.get("alt") or attrs.get("title") or "", m.group(0)))
    found.sort()
    out: list[ExtractedImage] = []
    for start, end, src, alt, markup in found:
        im = _load_markdown_src(src, base_dir)
        if im is None:
            continue
        # 上下文限制在图片所在的小节内：不跨过标题，避免相邻小节的内容"串"到这张图上
        before = text[max(0, start - cfg.context_chars):start]
        after = text[end:end + cfg.context_chars]
        heads = list(_MD_HEADING.finditer(before))
        if heads:
            before = before[heads[-1].start():]
        nxt = _MD_HEADING.search(after)
        if nxt:
            after = after[:nxt.start()]
        around = before + " " + after
        around = re.sub(r"^#{1,6}\s+", "", _HTML_IMAGE.sub(" ", _MD_IMAGE.sub(" ", around)), flags=re.M)
        out.append(ExtractedImage(image=im, sha1=pixel_sha1(im), origin="markdown", caption=alt.strip(), context=_clip(around, cfg.context_chars), markup=markup))
    return out


def _load_markdown_src(src: str, base_dir: Path | None) -> Any | None:
    if src.startswith("data:image"):
        try:
            return decode_image_b64(src)
        except ValueError:
            return None
    if re.match(r"^[a-z][a-z0-9+.-]*://", src, re.I):
        return None  # 远程图片不抓取
    if base_dir is None:
        return None
    from urllib.parse import unquote

    p = (base_dir / unquote(src.split("#")[0].split("?")[0])).resolve()
    if not p.is_file() or p.suffix.lower().lstrip(".") not in IMAGE_EXTENSIONS:
        return None
    return open_image(p.read_bytes())


def _from_mineru(doc: ParsedDocument, mineru_out_dir: Path | None) -> list[ExtractedImage]:
    """MinerU Content List 的 image / chart / table 块。MinerU 4.x 的 img_path 是内嵌的 data URI；
    旧版 / 落盘模式下是相对路径，到输出目录里找。"""
    roots = [Path(doc.metadata.source).parent]
    if mineru_out_dir is not None:
        roots = [mineru_out_dir / doc.doc_id, mineru_out_dir] + roots
    out: list[ExtractedImage] = []
    runs: list[int] = []                 # 每张图属于第几组"同页连续的图片块"（中间没有正文）
    run, last = -1, None
    blocks = doc.blocks
    for i, b in enumerate(blocks):
        if b.type in (BlockType.TEXT, BlockType.LIST, BlockType.TITLE) and (b.text or "").strip():
            last = None
        img_path = b.meta.get("img_path")
        if not img_path or b.type not in (BlockType.IMAGE, BlockType.TABLE):
            continue
        im = None
        if str(img_path).startswith("data:image"):
            try:
                im = decode_image_b64(img_path)
            except ValueError:
                im = None
        else:
            p = Path(img_path)
            file = p if p.is_absolute() and p.is_file() else next((r / img_path for r in roots if (r / img_path).is_file()), None)
            im = open_image(file.read_bytes()) if file else None
        if im is None:
            continue
        # 上下文：同一页上紧邻的正文块
        near = [x.text for x in blocks[max(0, i - 2):i] + blocks[i + 1:i + 2] if x.type in (BlockType.TEXT, BlockType.LIST, BlockType.TITLE) and x.page == b.page and x.text.strip()]
        caption = str(b.meta.get("caption") or "")
        if not caption:
            # 版面模型有时把图注识别成独立的正文块（standard 档实测）：在前后相邻块里找"图N …"
            for x in blocks[i + 1:i + 3] + blocks[max(0, i - 2):i][::-1]:
                if x.page == b.page and x.type in (BlockType.TEXT, BlockType.CAPTION) and _CAPTION.match(x.text or "") and len(x.text) <= 120:
                    caption = x.text.strip()
                    break
        # MinerU 的 image_analysis 已给出图片描述 / 图表数据：直接复用，不再调一次 VLM
        description = str(b.meta.get("description") or "") if b.type == BlockType.IMAGE else ""
        out.append(ExtractedImage(image=im, sha1=pixel_sha1(im), origin="mineru", page=b.page, bbox=b.bbox, caption=caption[:200], description=description, context=_clip("\n".join(near), 300)))
        if last != b.page:
            run += 1
        runs.append(run)
        last = b.page
    # 复合图被版面模型拆成几块时只有一块带图注（standard 档实测：CLIP 论文图 14 拆成 3 块、图 17 / 18 各 2 块）。
    # 同一组里恰好一个图注时，其余没有图注的块继承它——否则这些碎片既搜不到，也按图号找不到
    for r in set(runs):
        members = [im for im, k in zip(out, runs) if k == r]
        caps = {im.caption for im in members if im.caption}
        if len(members) > 1 and len(caps) == 1:
            for im in members:
                im.caption = im.caption or next(iter(caps))
    return out
